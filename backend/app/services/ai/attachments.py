from __future__ import annotations

import base64
import csv
import io
import json
from pathlib import Path
import re
import uuid
from typing import Any

from fastapi import HTTPException, UploadFile


UPLOAD_ROOT = Path("/app/data/ai_uploads")
MAX_FILE_SIZE = 12 * 1024 * 1024
ALLOWED_EXTENSIONS = {
    ".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md", ".xml",
    ".json", ".png", ".jpg", ".jpeg", ".webp",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
TEXT_EXTENSIONS = {".csv", ".txt", ".md", ".xml", ".json"}


def _safe_filename(name: str) -> str:
    name = Path(name or "arquivo").name
    name = re.sub(r"[^A-Za-z0-9._()\- áàâãéêíóôõúçÁÀÂÃÉÊÍÓÔÕÚÇ]", "_", name)
    return name[:180] or "arquivo"


def _extract_text(path: Path, extension: str) -> str:
    try:
        if extension in TEXT_EXTENSIONS:
            raw = path.read_bytes()
            for encoding in ("utf-8", "utf-8-sig", "latin-1"):
                try:
                    return raw.decode(encoding)
                except UnicodeDecodeError:
                    continue
            return raw.decode("utf-8", errors="replace")

        if extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return "\n\n".join((page.extract_text() or "") for page in reader.pages)

        if extension == ".docx":
            from docx import Document
            document = Document(str(path))
            return "\n".join(paragraph.text for paragraph in document.paragraphs)

        if extension == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(str(path), read_only=True, data_only=True)
            lines: list[str] = []
            for worksheet in workbook.worksheets:
                lines.append(f"# Planilha: {worksheet.title}")
                for row in worksheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values):
                        lines.append("\t".join(values))
            return "\n".join(lines)
    except Exception as exc:
        return f"[Não foi possível extrair o conteúdo textual automaticamente: {exc}]"

    return ""


async def save_attachment(file: UploadFile, user_id: int) -> dict[str, Any]:
    filename = _safe_filename(file.filename or "arquivo")
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                "Formato não suportado. Use PDF, DOCX, XLSX, CSV, TXT, MD, XML, "
                "JSON, PNG, JPG ou WEBP."
            ),
        )

    content = await file.read(MAX_FILE_SIZE + 1)
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="Arquivo maior que 12 MB.")

    attachment_id = str(uuid.uuid4())
    user_dir = UPLOAD_ROOT / str(user_id)
    user_dir.mkdir(parents=True, exist_ok=True)

    stored_name = f"{attachment_id}{extension}"
    path = user_dir / stored_name
    path.write_bytes(content)

    extracted_text = "" if extension in IMAGE_EXTENSIONS else _extract_text(path, extension)
    text_path = user_dir / f"{attachment_id}.txt"
    if extracted_text:
        text_path.write_text(extracted_text[:500_000], encoding="utf-8")

    metadata = {
        "id": attachment_id,
        "user_id": user_id,
        "filename": filename,
        "extension": extension,
        "content_type": file.content_type or "application/octet-stream",
        "size": len(content),
        "kind": "image" if extension in IMAGE_EXTENSIONS else "document",
        "stored_name": stored_name,
        "has_text": bool(extracted_text),
    }
    (user_dir / f"{attachment_id}.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return metadata


def load_attachments(attachment_ids: list[str], user_id: int) -> list[dict[str, Any]]:
    user_dir = UPLOAD_ROOT / str(user_id)
    results: list[dict[str, Any]] = []
    for attachment_id in attachment_ids[:8]:
        if not re.fullmatch(r"[0-9a-fA-F\-]{36}", attachment_id or ""):
            continue
        metadata_path = user_dir / f"{attachment_id}.json"
        if not metadata_path.exists():
            continue
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        if int(metadata.get("user_id", -1)) != int(user_id):
            continue

        text_path = user_dir / f"{attachment_id}.txt"
        metadata["text"] = (
            text_path.read_text(encoding="utf-8", errors="replace")[:120_000]
            if text_path.exists()
            else ""
        )
        stored_path = user_dir / str(metadata.get("stored_name"))
        if metadata.get("kind") == "image" and stored_path.exists():
            metadata["image_base64"] = base64.b64encode(stored_path.read_bytes()).decode("ascii")
        results.append(metadata)
    return results


def attachments_prompt_context(items: list[dict[str, Any]]) -> str:
    blocks: list[str] = []
    for item in items:
        name = item.get("filename", "arquivo")
        if item.get("kind") == "image":
            blocks.append(f"ARQUIVO ANEXADO: {name} (imagem enviada ao modelo para análise visual)")
            continue
        text = (item.get("text") or "").strip()
        if not text:
            blocks.append(f"ARQUIVO ANEXADO: {name} (sem texto extraível)")
            continue
        blocks.append(f"ARQUIVO ANEXADO: {name}\n---\n{text}\n---")
    return "\n\n".join(blocks)
